import io, re
from docx import Document

d = Document(r"C:\Users\20262\Desktop\烟头行为监测\.tmp-doc2\v2.docx")
parts = []
for p in d.paragraphs:
    if p.text.strip():
        parts.append(p.text.strip())
for t in d.tables:
    for row in t.rows:
        seen = []
        for c in row.cells:
            v = c.text.strip().replace("\n", " ")
            if not seen or seen[-1] != v:
                seen.append(v)
        parts.append(" | ".join(seen))
text = "\n".join(parts)
with open(r"C:\Users\20262\Desktop\烟头行为监测\.tmp-doc2\v2.txt", "w", encoding="utf-8") as f:
    f.write(text)
print("chars", len(text))

keys = ["捡", "拾", "扔", "抛掷", "ViFi", "CLIP", "导航", "姿态", "光流", "CBAM", "SE", "ByteTrack",
        "85%", "10%", "mAP", "召回", "误报", "证据链", "落地", "持烟", "吸烟", "垃圾桶", "排班",
        "热力", "YOLOv8", "YOLO", "边缘", "跨摄像头", "重识别", "标注", "数据集"]
for k in keys:
    print(f"{k}: {text.count(k)}")
